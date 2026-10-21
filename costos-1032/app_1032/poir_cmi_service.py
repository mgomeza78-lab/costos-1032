# -*- coding: utf-8 -*-
"""Ingresos por proyecto: CMI por activo y modificaciones del POIR (Res. CRA 1032/688).

El VP de los ingresos recuperados vía CMI es activo por activo
(VP(CMI688ⱼ×Qt), art. 2.1.2.1.3.5.5). El CMI de la 688 pudo modificarse durante su
aplicación (modificaciones del POIR), de modo que un activo puede tener distintos
CMI según el período. Este módulo gestiona:
  · las modificaciones del POIR (cuántas y en qué fechas; 0 = original);
  · el CMI de cada activo bajo cada modificación, cargado por Excel (uno por modificación);
  · el CMI aplicable a cada año tarifario según la modificación vigente.

El cálculo del VP(CMI×Q) por activo y su integración a la BCR₀ vive en bcr_service.
"""
import io
import unicodedata
from datetime import date
from sqlalchemy import text

# Columnas esperadas en el Excel (encabezados flexibles por alias normalizado)
ALIAS = {
    "codigo": ["codigo", "id", "idactivo", "proyecto", "codactivo", "codigoactivo"],
    "nombre": ["nombre", "activo", "descripcion", "nombreactivo"],
    "componente": ["componente", "comp"],
    "origen": ["origen"],
    "via688": ["via688", "via", "valordelactivo", "valoractivo", "vi", "via_688"],
    "prr": ["prr", "plazo", "plazoremanente", "plazorecuperacion", "vidautilremanente", "vur"],
    "ejecucion_real": ["ejecucion", "ejecucionreal", "valorejecutado", "ejecutado"],
    "anio_poir": ["aniopoir", "anio", "ano", "anooperacion", "aniooperacion", "anioentrada"],
    "cmi": ["cmi", "cmi688", "cmicobrado", "cmiaplicado"],
}


def _norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return "".join(ch for ch in s.lower() if ch.isalnum())


def _mapa_columnas(fila_encab):
    """Mapea índice de columna → nombre de campo, a partir de la fila de encabezados."""
    norm2field = {}
    for field, aliases in ALIAS.items():
        for a in aliases:
            norm2field[a] = field
    col = {}
    for i, c in enumerate(fila_encab):
        f = norm2field.get(_norm(c))
        if f and f not in col.values():
            col[i] = f
    return col


def _num(x):
    if x is None or x == "":
        return None
    try:
        return float(str(x).replace("$", "").replace(" ", "").replace(".", "").replace(",", ".")) \
            if isinstance(x, str) else float(x)
    except (TypeError, ValueError):
        return None


def leer_excel_cmi(raw):
    """Lee el Excel de una modificación → (filas, errores). Cada fila es un dict con los campos mapeados."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    except Exception as e:
        return [], [f"no se pudo abrir el Excel ({type(e).__name__})"]
    ws = wb[wb.sheetnames[0]]
    filas, err = [], []
    col = {}
    encab_ok = False
    for r in ws.iter_rows(values_only=True):
        if not encab_ok:
            m = _mapa_columnas(r)
            if "codigo" in m.values() and "cmi" in m.values():
                col = m
                encab_ok = True
            continue
        if r is None or all(c is None or str(c).strip() == "" for c in r):
            continue
        d = {}
        for i, field in col.items():
            v = r[i] if i < len(r) else None
            if field in ("codigo", "nombre", "origen"):
                d[field] = str(v).strip() if v is not None else None
            elif field == "componente":
                d[field] = int(_num(v)) if _num(v) is not None else None
            elif field == "anio_poir":
                d[field] = int(_num(v)) if _num(v) is not None else None
            else:
                d[field] = _num(v)
        if d.get("codigo"):
            filas.append(d)
    if not encab_ok:
        err.append("no se encontraron los encabezados (se requieren al menos 'codigo' y 'cmi')")
    return filas, err


def plantilla_bytes():
    """Genera la plantilla Excel para cargar el CMI por activo de una modificación."""
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CMI_por_activo"
    encab = ["codigo", "nombre", "componente", "origen", "via688", "prr", "ejecucion_real", "anio_poir", "cmi"]
    ws.append(encab)
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.append(["AC-001", "Red matriz sector norte", 1, "VI287", 41200000000, 22, "", "", 782.69])
    ws.append(["AC-POIR-012", "Ampliación conducción", 1, "POIR688", "", 40, 300000000, 5, 120.50])
    nota = wb.create_sheet("INSTRUCCIONES")
    nota["A1"] = ("Una fila por activo/proyecto. 'cmi' es el CMI688 ($/m³) del activo BAJO ESTA modificación. "
                  "Suba un archivo por cada modificación del POIR (incluida la 0=original). "
                  "componente: 1 infraestructura, 2 ambiental/riesgo, 3 admin/operativo. "
                  "Para activos del POIR puede dejar 'via688' vacío y cargar 'ejecucion_real' y 'anio_poir': "
                  "el sistema calcula VIA = ejecucion/(1+r688)^anio_poir.")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------- Modificaciones del POIR ----------------
def mods_list(conn, empresa_id, servicio):
    return conn.execute(text(
        "SELECT numero, fecha, descripcion FROM poir_modificacion "
        "WHERE empresa_id=:e AND servicio=:s ORDER BY numero"),
        {"e": empresa_id, "s": servicio}).mappings().all()


def mods_set(conn, empresa_id, servicio, fechas, fecha_inicio_688=None):
    """Fija el nº de modificaciones y sus fechas. `fechas` = lista de fechas de las modificaciones 1..N.
    La 0 (POIR original) toma la fecha de inicio de la 688."""
    conn.execute(text("DELETE FROM poir_modificacion WHERE empresa_id=:e AND servicio=:s"),
                 {"e": empresa_id, "s": servicio})
    conn.execute(text(
        "INSERT INTO poir_modificacion(empresa_id,servicio,numero,fecha,descripcion) "
        "VALUES (:e,:s,0,:f,'POIR original')"),
        {"e": empresa_id, "s": servicio, "f": fecha_inicio_688 or None})
    n = 0
    for k, f in enumerate(fechas or [], start=1):
        conn.execute(text(
            "INSERT INTO poir_modificacion(empresa_id,servicio,numero,fecha,descripcion) "
            "VALUES (:e,:s,:n,:f,:d)"),
            {"e": empresa_id, "s": servicio, "n": k, "f": (f or None),
             "d": f"Modificación {k}"})
        n += 1
    return n


# ---------------- CMI por activo ----------------
def cmi_por_codigo(conn, empresa_id, servicio):
    """{codigo: {modificacion: cmi}}."""
    out = {}
    for r in conn.execute(text(
        "SELECT codigo, modificacion, cmi FROM activo_cmi WHERE empresa_id=:e AND servicio=:s"),
            {"e": empresa_id, "s": servicio}).mappings().all():
        out.setdefault(r["codigo"], {})[int(r["modificacion"])] = \
            float(r["cmi"]) if r["cmi"] is not None else None
    return out


def ingesta_modificacion(conn, empresa_id, servicio, modificacion, filas, usuario_id):
    """Carga el Excel de una modificación: upsert de activos (por código) y del CMI de la modificación."""
    n_act = n_cmi = 0
    for f in filas:
        cod = (f.get("codigo") or "").strip()
        if not cod:
            continue
        row = conn.execute(text(
            "SELECT id FROM bcr_activo WHERE empresa_id=:e AND servicio=:s AND codigo=:c"),
            {"e": empresa_id, "s": servicio, "c": cod}).scalar()
        comp = f.get("componente") if f.get("componente") in (1, 2, 3) else None
        attrs = {"e": empresa_id, "s": servicio, "c": cod,
                 "n": f.get("nombre") or cod, "cmp": comp or 1,
                 "og": f.get("origen") or "POIR688", "via": f.get("via688"),
                 "pr": f.get("prr"), "er": f.get("ejecucion_real"), "ap": f.get("anio_poir"),
                 "u": usuario_id}
        if row:
            conn.execute(text(
                "UPDATE bcr_activo SET nombre=COALESCE(:n,nombre), "
                " componente=COALESCE(:cmp,componente), origen=COALESCE(:og,origen), "
                " via688=COALESCE(:via,via688), plazo_remanente=COALESCE(:pr,plazo_remanente), "
                " ejecucion_real=COALESCE(:er,ejecucion_real), anio_poir=COALESCE(:ap,anio_poir), "
                " actualizado_en=now() WHERE empresa_id=:e AND servicio=:s AND codigo=:c"), attrs)
        else:
            conn.execute(text(
                "INSERT INTO bcr_activo(empresa_id,servicio,componente,nombre,codigo,origen,"
                " via688,plazo_remanente,ejecucion_real,anio_poir,cargado_por) "
                "VALUES (:e,:s,:cmp,:n,:c,:og,:via,:pr,:er,:ap,:u)"), attrs)
            n_act += 1
        if f.get("cmi") is not None:
            conn.execute(text(
                "INSERT INTO activo_cmi(empresa_id,servicio,codigo,modificacion,cmi) "
                "VALUES (:e,:s,:c,:m,:cmi) ON CONFLICT (empresa_id,servicio,codigo,modificacion) "
                "DO UPDATE SET cmi=excluded.cmi"),
                {"e": empresa_id, "s": servicio, "c": cod, "m": modificacion, "cmi": f.get("cmi")})
            n_cmi += 1
    return n_act, n_cmi


def cmi_aplicable(mods, cmi_de_codigo, year_start):
    """CMI del activo en un año que inicia en `year_start`, según la modificación vigente."""
    best = 0
    for m in mods:
        if m["fecha"] and year_start and m["fecha"] <= year_start and int(m["numero"]) > best:
            best = int(m["numero"])
    for k in range(best, -1, -1):
        if k in cmi_de_codigo and cmi_de_codigo[k] is not None:
            return cmi_de_codigo[k]
    return None


def resumen_cmi(conn, empresa_id, servicio):
    """Por activo: su CMI por modificación (para la tabla de la UI)."""
    cc = cmi_por_codigo(conn, empresa_id, servicio)
    nombres = dict(conn.execute(text(
        "SELECT codigo, nombre FROM bcr_activo WHERE empresa_id=:e AND servicio=:s AND codigo IS NOT NULL"),
        {"e": empresa_id, "s": servicio}).all())
    return [{"codigo": c, "nombre": nombres.get(c, c), "cmi": cc[c]} for c in sorted(cc)]
